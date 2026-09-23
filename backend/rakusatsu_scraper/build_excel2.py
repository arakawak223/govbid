#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""スクレイピング結果＋PR広報リスト照合分を統合し、落札企業Excel＋0616差分を生成。"""
import json, csv, re, glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
SP = '/tmp/claude-1000/-workspaces-govbid/ef35a61d-f5dd-4eb5-9dfb-6338a9ff2e0d/scratchpad'
ROOT = '/workspaces/govbid'

def norm(s):
    return re.sub(r'[\s　、。・,\.\-ー（）\(\)「」『』【】／/！!？?]+', '', str(s or ''))
def core(title):
    t = re.sub(r'^令和[０-９\d]+年度|^令和[０-９\d]+～[０-９\d]+年度|^R[０-９\d]+年度|^【.*?】', '', str(title))
    t = re.sub(r'（.*?）|\(.*?\)', '', t)
    t = re.sub(r'(に係る企画提案競技.*$|業務委託$|委託業務$|業務$|委託$|事業$|等$)', '', t)
    return norm(t)
def cwin(s):  # 企業名正規化（支店・営業所差を吸収）
    return re.sub(r'(株式会社|有限会社|合同会社).*?$', r'\1', '') or norm(re.sub(r'(福岡支店|北九州支店|北九州支社|宮崎営業所|長崎支店|.*支店|.*支社|.*営業所)$','',str(s or '')))

# --- スクレイピング結果 ---
res = [x for x in json.load(open(f'{SP}/scrape_v3.json')) if x['winner']]
for x in res:
    if x['winner'] == '株式会社J T' and 'JTB' in norm(x['evidence']):
        x['winner'] = '株式会社JTB福岡支店'
    x['winner'] = re.sub(r'\b([A-Z]) (?=[A-Z]\b)', r'\1', x['winner'])
tg = {r['row']: r for r in csv.DictReader(open(f'{SP}/targets.csv', encoding='utf-8'))}
recs = []
for x in sorted(res, key=lambda r: int(r['row'])):
    t = tg.get(str(x['row']), {})
    recs.append({'出典': 'スクレイピング', '調査日': t.get('調査日',''), '都道府県': x['都道府県'],
                 '管轄': t.get('管轄',''), '案件名': x['イベント名'], '金額': int(x['委託料上限']),
                 'winner': x['winner'], 'label': x['label'], 'url': x['result_url'] or x['URL'],
                 'ev': x['evidence'][:120]})
# scraper内 重複除去
seen, uniq = set(), []
for r in recs:
    k = (core(r['案件名']), norm(r['winner']))
    if k in seen: continue
    seen.add(k); uniq.append(r)

# --- PR広報リスト照合分 ---
pr = json.load(open(f'{SP}/pr_winners.json'))
existing_win = {norm(r['winner']) for r in uniq}
existing_core = {core(r['案件名']) for r in uniq}
def cbase(w):  # 企業名の基幹（支店/支社/営業所を除去）
    return norm(re.sub(r'(北九州支店|北九州支社|福岡支店|福岡支社|宮崎営業所|長崎支店|.{0,4}支店|.{0,4}支社|.{0,4}営業所)$', '', str(w or '')))
def share_run(a, b, n=5):  # a,b がn文字以上の連続部分文字列を共有するか
    a, b = norm(a), norm(b)
    return any(a[i:i+n] in b for i in range(0, max(1, len(a)-n+1)))
def already(p):
    nw, wb = norm(p['受託者']), cbase(p['受託者'])
    for r in uniq:
        rw, rb = norm(r['winner']), cbase(r['winner'])
        same_co = (wb and rb and (wb[:6] in rb or rb[:6] in wb)) or nw[:6] == rw[:6]
        if same_co and share_run(p['案件名'], r['案件名'], 5):
            return True
    return False
def prdate(d):
    m = re.match(r'(\d+)/(\d+)', str(d))
    return f'2026-{int(m.group(1)):02d}-{int(m.group(2)):02d}' if m else str(d)
added = 0
for p in pr:
    if already(p):
        continue
    recs_new = {'出典': 'PR広報リスト照合', '調査日': prdate(p['date']), '都道府県': '福岡県',
                '管轄': p['発注主体'], '案件名': p['案件名'],
                '金額': int(p['金額']) if str(p['金額']).isdigit() else '', 'winner': p['受託者'],
                'label': f"PRリスト({p['リンク種別']})", 'url': p['リンク'], 'ev': p['備考'] or 'PR広報リストで確認済み'}
    uniq.append(recs_new); added += 1
print(f'スクレイピング {len(seen)}種 + PR照合追加 {added}種 = 計 {len(uniq)}件')

# --- 0616既公表分 ---
f0616 = glob.glob(f'{ROOT}/*0616*.xlsx')[0]
wb0 = openpyxl.load_workbook(f0616, read_only=True)
prev = set()
for sn, ci in [('抽出落札企業一覧', 2), ('落札・採択企業リスト（広報PR）', 2)]:
    for r in wb0[sn].iter_rows(min_row=2, values_only=True):
        if r[ci]: prev.add(core(r[ci]))
wb0.close()
for r in uniq:
    r['新規'] = core(r['案件名']) not in prev
new_only = [r for r in uniq if r['新規']]
print(f'0616既公表 {len(prev)}種 / 新規公表 {len(new_only)} / 全 {len(uniq)}')

# --- Excel ---
HDR = ['No', '出典', '調査日/公表', '都道府県', '発注主体/管轄', '案件名', '金額(円)', '落札企業',
       '判定ラベル', '新規', '結果URL', '抜粋(根拠)']
thin = Side(style='thin', color='CCCCCC'); border = Border(thin, thin, thin, thin)
hdr_fill = PatternFill('solid', fgColor='1F4E78'); new_fill = PatternFill('solid', fgColor='FFF2CC')
pr_fill = PatternFill('solid', fgColor='E2EFDA')
def write_sheet(ws, rows):
    ws.append(HDR)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr_fill
        c.alignment = Alignment('center', 'center', wrap_text=True); c.border = border
    for i, x in enumerate(rows, 1):
        ws.append([i, x['出典'], x['調査日'], x['都道府県'], x['管轄'], x['案件名'],
                   x['金額'], x['winner'], x['label'], '★新規' if x['新規'] else '既出', x['url'], x['ev']])
        r = ws.max_row
        for c in ws[r]: c.border = border
        if isinstance(x['金額'], int): ws.cell(r, 7).number_format = '#,##0'
        if x['新規']:
            for c in ws[r]: c.fill = new_fill
        elif x['出典'].startswith('PR'):
            for c in ws[r]: c.fill = pr_fill
    for i, w in enumerate([5,15,12,8,12,42,13,34,18,8,52,50], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
wb = openpyxl.Workbook()
write_sheet(wb.active, uniq); wb.active.title = '落札企業リスト_0721'
write_sheet(wb.create_sheet('新規公表分_0616差分'), new_only)
out = f'{ROOT}/九州都道府県別プロポーザルまとめ_0721_落札企業反映.xlsx'
wb.save(out)
print('saved:', out)
print('\n--- PR照合で追加した分 ---')
for r in uniq:
    if r['出典'].startswith('PR'):
        print(f"  {r['管轄']}|{r['案件名'][:28]} → {r['winner']}  [{'新規' if r['新規'] else '既出'}]")
